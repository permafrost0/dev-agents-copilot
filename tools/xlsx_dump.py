"""詳細設計書（xlsx）の指定範囲だけをテキストに変換する。

- 指示された シート × セル範囲 だけを出力する。範囲外のセルの値は一切出力しない
- 出力するもの：セルの値、範囲に掛かる結合セル、セルのコメント、アンカーが範囲内にある図形の文字
- 画像は扱わない（存在の報告もしない）
- ファイルはバイト列として読み込むだけで、保存処理は持たない（設計書を変更しない）

使い方:
  python tools/xlsx_dump.py --feature <機能名>
      work/<機能名>/00_instruction.md の『対象範囲』を全て抽出し、
      work/<機能名>/_ai/extract/ に1範囲1ファイルで出力する
  python tools/xlsx_dump.py --book <ファイル名> --sheet <シート名> --range A1:C10 [--range ...]
      標準出力に出力する（ファイル名は input/設計書/ 直下から探す）
"""
from __future__ import annotations

import argparse
import datetime as dt
import io
import posixpath
import re
import sys
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import _common as C  # noqa: E402

NS = {
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}


class DumpError(Exception):
    pass


def resolve_book(name: str) -> Path:
    """設計書は input/設計書/ の直下に置く決まり。パス指定は受け付けない。"""
    if "/" in name or "\\" in name:
        raise DumpError(f"設計書はファイル名だけで指定してください（input/設計書/ 直下を探します）: {name}")
    p = C.design_doc_dir() / name
    if not p.is_file():
        raise DumpError(f"設計書が見つかりません: input/設計書/{name}")
    return p


def fmt_value(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (dt.datetime, dt.date, dt.time)):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v)
    return s.replace("\\", "\\\\").replace("|", "\\|").replace("\r\n", "\n").replace("\n", "<br>")


# ---------------------------------------------------------------- 図形の文字
def _sheet_xml_path(z: zipfile.ZipFile, sheet: str) -> str | None:
    wb = ET.fromstring(z.read("xl/workbook.xml"))
    rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    rid = None
    for s in wb.iter(f"{{{NS['main']}}}sheet"):
        if s.get("name") == sheet:
            rid = s.get(f"{{{NS['r']}}}id")
            break
    if rid is None:
        return None
    for r in rels.iter(f"{{{NS['rel']}}}Relationship"):
        if r.get("Id") == rid:
            target = r.get("Target")
            if target.startswith("/"):
                return target.lstrip("/")
            return posixpath.normpath(posixpath.join("xl", target))
    return None


def _drawing_paths(z: zipfile.ZipFile, sheet_xml: str) -> list[str]:
    d, f = posixpath.split(sheet_xml)
    rels_path = posixpath.join(d, "_rels", f + ".rels")
    if rels_path not in z.namelist():
        return []
    rels = ET.fromstring(z.read(rels_path))
    out = []
    for r in rels.iter(f"{{{NS['rel']}}}Relationship"):
        if r.get("Type", "").endswith("/drawing"):
            t = r.get("Target")
            out.append(t.lstrip("/") if t.startswith("/") else posixpath.normpath(posixpath.join(d, t)))
    return out


def _shape_text(sp: ET.Element) -> str:
    paras = []
    for p in sp.iter(f"{{{NS['a']}}}p"):
        t = "".join(x.text or "" for x in p.iter(f"{{{NS['a']}}}t"))
        if t:
            paras.append(t)
    return "\n".join(paras)


def shapes_in_range(data: bytes, sheet: str, bounds: tuple) -> list[dict]:
    """アンカー（左上）が範囲内にある図形の文字を返す。画像・グラフは対象外。"""
    out = []
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        sx = _sheet_xml_path(z, sheet)
        if not sx:
            return out
        for dp in _drawing_paths(z, sx):
            if dp not in z.namelist():
                continue
            root = ET.fromstring(z.read(dp))
            for anchor in list(root):
                tag = anchor.tag.split("}")[-1]
                if tag not in ("twoCellAnchor", "oneCellAnchor"):
                    continue  # absoluteAnchor はセル位置を持たないので対象外
                frm = anchor.find("xdr:from", NS)
                if frm is None:
                    continue
                col = int(frm.find("xdr:col", NS).text) + 1
                row = int(frm.find("xdr:row", NS).text) + 1
                if not (bounds[0] <= row <= bounds[2] and bounds[1] <= col <= bounds[3]):
                    continue
                # 図形（sp）とグループ内の図形だけを見る。pic（画像）と graphicFrame（グラフ）は見ない
                shapes = []
                for child in anchor:
                    ctag = child.tag.split("}")[-1]
                    if ctag == "sp":
                        shapes.append(child)
                    elif ctag == "grpSp":
                        shapes.extend(child.iter(f"{{{NS['xdr']}}}sp"))
                for sp in shapes:
                    text = _shape_text(sp)
                    if not text:
                        continue
                    nv = sp.find("xdr:nvSpPr/xdr:cNvPr", NS)
                    out.append(
                        {
                            "anchor": f"{C.num_to_col(col)}{row}",
                            "name": nv.get("name", "") if nv is not None else "",
                            "text": text,
                        }
                    )
    return out


# ---------------------------------------------------------------- 本体
def dump(book: str, sheet: str, rng: str) -> str:
    from openpyxl import load_workbook
    from openpyxl.cell.cell import MergedCell

    path = resolve_book(book)
    data = path.read_bytes()  # 読み取り専用で読み込み、以後ファイルには触れない
    bounds = C.parse_range(rng)
    wb = load_workbook(io.BytesIO(data), data_only=True)
    if sheet not in wb.sheetnames:
        raise DumpError(f"シートが見つかりません: {sheet}（{book}）")
    ws = wb[sheet]
    wb_f = None  # 数式を表示するために必要になった時だけ読む

    ref_fmt = f"[設計書: {book}/{sheet}!<セル>]"
    lines = [
        "# 抽出結果",
        f"- 設計書: {book}",
        f"- シート: {sheet}",
        f"- 範囲: {rng.upper()}",
        f"- 参照の書き方: `{ref_fmt}`",
        "- 注意: この範囲の外は抽出していない。範囲外を参照しないと読めない記述は質問（Q）にすること",
        "",
        "## セル",
        "| セル | 値 |",
        "|---|---|",
    ]
    count = 0
    for row in ws.iter_rows(
        min_row=bounds[0], min_col=bounds[1], max_row=bounds[2], max_col=bounds[3]
    ):
        for cell in row:
            if isinstance(cell, MergedCell):
                continue  # 結合セルの内側は左上で代表する
            v = cell.value
            if v is None:
                if wb_f is None:
                    wb_f = load_workbook(io.BytesIO(data), data_only=False)
                fv = wb_f[sheet][cell.coordinate].value
                if isinstance(fv, str) and fv.startswith("="):
                    v = f"（数式・計算結果なし）{fv}"
            s = fmt_value(v)
            if s == "":
                continue
            lines.append(f"| {cell.coordinate} | {s} |")
            count += 1
    if count == 0:
        lines.append("| （値のあるセルなし） | |")

    merged = []
    for mr in ws.merged_cells.ranges:
        mb = (mr.min_row, mr.min_col, mr.max_row, mr.max_col)
        if not C.ranges_intersect(mb, bounds):
            continue
        tl = f"{C.num_to_col(mr.min_col)}{mr.min_row}"
        inside_tl = bounds[0] <= mr.min_row <= bounds[2] and bounds[1] <= mr.min_col <= bounds[3]
        crosses = not C.range_within(mr.coord, rng)
        note = f"左上 {tl}"
        if crosses:
            note += "、指定範囲の外にまたがる"
        if not inside_tl:
            note += "。左上セルが範囲外のため値は出力しない（必要なら質問にすること）"
        merged.append(f"- {mr.coord}（{note}）")
    lines += ["", "## 結合セル"] + (merged or ["- なし"])

    comments = []
    for row in ws.iter_rows(
        min_row=bounds[0], min_col=bounds[1], max_row=bounds[2], max_col=bounds[3]
    ):
        for cell in row:
            cm = getattr(cell, "comment", None)
            if cm is not None and cm.text:
                comments.append(f"- {cell.coordinate}: {fmt_value(cm.text.strip())}")
    lines += ["", "## コメント"] + (comments or ["- なし"])

    shapes = [
        f"- {s['anchor']}（{s['name']}）: {fmt_value(s['text'])}"
        for s in shapes_in_range(data, sheet, bounds)
    ]
    lines += ["", "## 図形内の文字（アンカーが範囲内のもの）"] + (shapes or ["- なし"])
    return "\n".join(lines) + "\n"


def safe_name(s: str) -> str:
    return re.sub(r'[\\/:*?"<>|\s]+', "_", s)


def dump_feature(feature: str) -> list[Path]:
    rows = C.load_instruction(feature)
    if not rows:
        raise DumpError(f"work/{feature}/00_instruction.md の『対象範囲』表に行がありません")
    out_dir = C.feature_dir(feature) / "_ai" / "extract"
    out_dir.mkdir(parents=True, exist_ok=True)
    written = []
    for i, r in enumerate(rows, 1):
        text = dump(r["book"], r["sheet"], r["range"])
        if r.get("note"):
            text = text.replace("\n## セル", f"\n- 指示書の補足: {r['note']}\n\n## セル", 1)
        p = out_dir / f"{i:02d}_{safe_name(r['sheet'])}_{r['range'].replace(':', '-')}.md"
        p.write_text(text, encoding="utf-8")
        written.append(p)
    return written


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--feature")
    ap.add_argument("--book")
    ap.add_argument("--sheet")
    ap.add_argument("--range", action="append", dest="ranges")
    a = ap.parse_args(argv)
    try:
        if a.feature:
            for p in dump_feature(a.feature):
                print(f"出力: {p.relative_to(C.root()).as_posix()}")
            return 0
        if not (a.book and a.sheet and a.ranges):
            ap.error("--feature か、--book/--sheet/--range の組を指定してください")
        for rng in a.ranges:
            print(dump(a.book, a.sheet, rng))
        return 0
    except DumpError as e:
        print(f"エラー: {e}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
